"""Exportação do relatório de demandas em PDF, Excel (.xlsx) e CSV.

Os três formatos recebem o mesmo dicionário `relatorio`, montado em
app.montar_relatorio():

    empresa       nome da empresa (cabeçalho)
    titulo        título do relatório
    gerado_em     datetime da geração (cabeçalho)
    filtros_texto "Prioridade Alta · Jan–Mar 2025" (rodapé)
    dados         resultado de app.calcular_painel() com os mesmos filtros
    demandas      lista de dicts, uma por demanda do recorte filtrado

Todo relatório tem: cabeçalho com nome da empresa e data de geração, e
rodapé com os filtros aplicados.
"""
import csv
import io
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

COR_ESCURA = '#2b2b2b'
COR_ATRASADA = '#b91c1c'
COR_ZEBRA = '#f3f4f6'

FORMATOS = {
    'pdf': ('application/pdf', 'pdf'),
    'xlsx': ('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'xlsx'),
    'csv': ('text/csv; charset=utf-8', 'csv'),
}


# ---------------------------------------------------------------- utilidades

def _num(valor):
    """3.3 -> '3,3'."""
    return str(valor).replace('.', ',')


def _data(valor, com_hora=False):
    if valor is None:
        return ''
    return valor.strftime('%d/%m/%Y %H:%M' if com_hora else '%d/%m/%Y')


def _gerado_em(rel):
    return rel['gerado_em'].strftime('%d/%m/%Y às %H:%M')


def _latin(texto):
    """A fonte padrão do PDF (Helvetica) cobre o alfabeto latino com
    acentos. Caracteres fora dele viram '?', em vez de quadrados pretos."""
    return str(texto).encode('cp1252', errors='replace').decode('cp1252')


def _linhas_resumo(dados):
    """(indicador, valor, observação) dos KPIs, iguais aos do painel."""
    k = dados['kpis']
    if dados['tempo_ponderado'] is not None:
        tempo = f"{_num(dados['tempo_ponderado'])} dias"
        tempo_obs = (f"Média ponderada: críticas (prioridade Alta) pesam {dados['peso_critica']}x. "
                     f"Base: {dados['tempo_base']} concluída(s). Canceladas ficam de fora.")
        simples = f"{_num(dados['tempo_simples'])} dias"
    else:
        tempo, tempo_obs, simples = '—', 'Nenhuma demanda concluída neste recorte.', '—'
    return [
        ('Total de demandas', k['total'], ''),
        ('Abertas', k['abertas'], f"{_num(k['pct_abertas'])}% do total ({k['em_andamento']} em andamento)"),
        ('Concluídas', k['concluidas'], f"{_num(k['pct_concluidas'])}% do total"),
        ('Atrasadas', k['atrasadas'], f"{_num(k['pct_atrasadas'])}% do total (abertas com prazo vencido)"),
        ('Canceladas', k['canceladas'], f"{_num(k['pct_canceladas'])}% do total"),
        ('Tempo médio de resolução', tempo, tempo_obs),
        ('Tempo médio simples', simples, 'Sem ponderação, só para comparação'),
    ]


# ---------------------------------------------------------------------- CSV

def _csv_seguro(valor):
    """Evita 'injeção de fórmula': texto digitado por usuários que começa
    com = + - @ seria executado como fórmula ao abrir no Excel."""
    if isinstance(valor, str) and valor[:1] in ('=', '+', '-', '@', '\t', '\r'):
        return "'" + valor
    return valor


def gerar_csv(rel):
    """CSV da lista de demandas do recorte. Separador ';' e UTF-8 com BOM,
    que é o que o Excel em português abre corretamente."""
    saida = io.StringIO()
    w = csv.writer(saida, delimiter=';', lineterminator='\r\n')

    w.writerow([rel['empresa']])
    w.writerow([rel['titulo']])
    w.writerow([f"Gerado em {_gerado_em(rel)}"])
    w.writerow([])
    w.writerow(['Nº', 'Título', 'Solicitante', 'Responsável', 'Prioridade', 'Status',
                'Criação', 'Prazo', 'Conclusão', 'Atrasada'])
    for d in rel['demandas']:
        w.writerow([_csv_seguro(x) for x in (
            d['id'], d['titulo'], d['solicitante'], d['responsavel'], d['prioridade'],
            d['status'], _data(d['criacao'], True), _data(d['prazo']),
            _data(d['conclusao'], True), 'Sim' if d['atrasada'] else 'Não',
        )])
    w.writerow([])
    w.writerow([f"Gerado com: {rel['filtros_texto']}"])
    return ('\ufeff' + saida.getvalue()).encode('utf-8')


# -------------------------------------------------------------------- Excel

def _texto(ws, linha, coluna, valor, **estilo):
    """Escreve uma célula; texto nunca é interpretado como fórmula."""
    celula = ws.cell(row=linha, column=coluna, value=valor)
    if isinstance(valor, str):
        celula.data_type = 's'
    if 'font' in estilo:
        celula.font = estilo['font']
    if 'fill' in estilo:
        celula.fill = estilo['fill']
    if 'alinhamento' in estilo:
        celula.alignment = estilo['alinhamento']
    return celula


def _cabecalho_xlsx(ws, rel):
    _texto(ws, 1, 1, rel['empresa'], font=Font(bold=True, size=16))
    _texto(ws, 2, 1, rel['titulo'], font=Font(bold=True, size=12))
    _texto(ws, 3, 1, f"Gerado em {_gerado_em(rel)}", font=Font(italic=True, color='555555'))


def _rodape_xlsx(ws, linha, rel):
    _texto(ws, linha, 1, f"Gerado com: {rel['filtros_texto']}", font=Font(italic=True, color='555555'))


def _impressao_xlsx(ws, rel, linha_titulos=None):
    """Prepara a planilha para sair no papel: paisagem, cabeçalho e rodapé
    de impressão (empresa/data e filtros) e largura ajustada à folha."""
    def esc(t):  # '&' é caractere especial em cabeçalho/rodapé de impressão
        return str(t).replace('&', '&&')
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = esc(rel['empresa'])
    ws.oddHeader.right.text = esc('Gerado em ' + _gerado_em(rel))
    ws.oddFooter.left.text = esc('Gerado com: ' + rel['filtros_texto'])
    ws.oddFooter.right.text = 'Página &P de &N'
    if linha_titulos:
        ws.print_title_rows = f'{linha_titulos}:{linha_titulos}'


def _tabela_xlsx(ws, linha, colunas, linhas, larguras=None, vermelhas=()):
    """Escreve cabeçalho + linhas a partir de `linha`. Devolve a próxima
    linha livre. `vermelhas` = índices (0-based) das linhas em vermelho."""
    cinza = PatternFill('solid', fgColor='2B2B2B')
    for j, nome in enumerate(colunas, start=1):
        _texto(ws, linha, j, nome, font=Font(bold=True, color='FFFFFF'), fill=cinza,
               alinhamento=Alignment(vertical='center'))
    for i, valores in enumerate(linhas):
        for j, v in enumerate(valores, start=1):
            c = _texto(ws, linha + 1 + i, j, v, alinhamento=Alignment(vertical='top', wrap_text=True))
            if i in vermelhas:
                c.font = Font(color='B91C1C', bold=(j == 1))
            if hasattr(v, 'year'):
                c.number_format = 'DD/MM/YYYY HH:MM' if hasattr(v, 'hour') else 'DD/MM/YYYY'
                c.alignment = Alignment(horizontal='center', vertical='top')
    if larguras:
        for j, larg in enumerate(larguras, start=1):
            atual = ws.column_dimensions[get_column_letter(j)].width or 0
            ws.column_dimensions[get_column_letter(j)].width = max(atual, larg)
    return linha + 1 + len(linhas)


def gerar_xlsx(rel):
    dados = rel['dados']
    wb = Workbook()

    # --- Resumo
    ws = wb.active
    ws.title = 'Resumo'
    _cabecalho_xlsx(ws, rel)
    linha = _tabela_xlsx(ws, 5, ['Indicador', 'Valor', 'Observação'], _linhas_resumo(dados), [34, 16, 80])
    atrasada = [i for i, l in enumerate(_linhas_resumo(dados)) if l[0] == 'Atrasadas']
    for i in atrasada:
        ws.cell(row=6 + i, column=1).font = Font(color='B91C1C', bold=True)
        ws.cell(row=6 + i, column=2).font = Font(color='B91C1C', bold=True)
    linha = _tabela_xlsx(ws, linha + 1, ['Status', 'Demandas'],
                         [(x['nome'], x['qtd']) for x in dados['por_status']])
    linha = _tabela_xlsx(ws, linha + 1, ['Prioridade', 'Demandas'],
                         [(x['nome'], x['qtd']) for x in dados['por_prioridade']])
    _rodape_xlsx(ws, linha + 1, rel)
    _impressao_xlsx(ws, rel)

    # --- Demandas
    ws = wb.create_sheet('Demandas')
    _cabecalho_xlsx(ws, rel)
    linhas = [(d['id'], d['titulo'], d['solicitante'], d['responsavel'], d['prioridade'], d['status'],
               d['criacao'], d['prazo'], d['conclusao'], 'Sim' if d['atrasada'] else 'Não')
              for d in rel['demandas']]
    vermelhas = [i for i, d in enumerate(rel['demandas']) if d['atrasada']]
    fim = _tabela_xlsx(
        ws, 5,
        ['Nº', 'Título', 'Solicitante', 'Responsável', 'Prioridade', 'Status', 'Criação', 'Prazo', 'Conclusão', 'Atrasada'],
        linhas, [7, 46, 18, 18, 12, 14, 18, 13, 18, 10], vermelhas)
    if linhas:
        ws.auto_filter.ref = f'A5:J{5 + len(linhas)}'
    ws.freeze_panes = 'A6'
    _rodape_xlsx(ws, fim + 1, rel)
    _impressao_xlsx(ws, rel, linha_titulos=5)

    # --- Críticas
    ws = wb.create_sheet('Críticas')
    _cabecalho_xlsx(ws, rel)
    linhas = [(c['id'], c['titulo'], c['responsavel'], c['status'], c['prazo'],
               'Sim' if c['atrasada'] else 'Não') for c in dados['criticas_lista']]
    vermelhas = [i for i, c in enumerate(dados['criticas_lista']) if c['atrasada']]
    fim = _tabela_xlsx(ws, 5, ['Nº', 'Título', 'Responsável', 'Status', 'Prazo', 'Atrasada'],
                       linhas, [7, 46, 20, 14, 13, 10], vermelhas)
    _rodape_xlsx(ws, fim + 1, rel)
    _impressao_xlsx(ws, rel, linha_titulos=5)

    # --- Por responsável
    ws = wb.create_sheet('Por responsável')
    _cabecalho_xlsx(ws, rel)
    fim = _tabela_xlsx(ws, 5, ['Responsável', 'Demandas em aberto'],
                       [(p['nome'], p['qtd']) for p in dados['por_responsavel']], [30, 22])
    _rodape_xlsx(ws, fim + 1, rel)
    _impressao_xlsx(ws, rel, linha_titulos=5)

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


# ---------------------------------------------------------------------- PDF

class _CanvasNumerado(canvas.Canvas):
    """Canvas em duas passadas para desenhar 'Página X de Y' e repetir
    cabeçalho e rodapé em todas as páginas."""

    def __init__(self, *args, relatorio=None, **kwargs):
        super().__init__(*args, **kwargs)
        self._rel = relatorio
        self._paginas = []

    def showPage(self):
        self._paginas.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._paginas)
        for estado in self._paginas:
            self.__dict__.update(estado)
            self._moldura(total)
            super().showPage()
        super().save()

    def _moldura(self, total):
        rel = self._rel
        largura, altura = self._pagesize
        margem = 1.5 * cm

        # Cabeçalho: empresa à esquerda, data de geração à direita
        self.setFillColor(colors.HexColor(COR_ESCURA))
        self.setFont('Helvetica-Bold', 15)
        self.drawString(margem, altura - 1.2 * cm, _latin(rel['empresa']))
        self.setFont('Helvetica', 10)
        self.setFillColor(colors.HexColor('#555555'))
        self.drawRightString(largura - margem, altura - 1.2 * cm, _latin('Gerado em ' + _gerado_em(rel)))
        self.setFont('Helvetica', 9)
        self.drawString(margem, altura - 1.7 * cm, _latin(rel['titulo']))
        self.setStrokeColor(colors.HexColor('#cccccc'))
        self.line(margem, altura - 1.95 * cm, largura - margem, altura - 1.95 * cm)

        # Rodapé: filtros aplicados à esquerda, página à direita
        self.line(margem, 1.75 * cm, largura - margem, 1.75 * cm)
        self.setFont('Helvetica', 8.5)
        texto = _latin('Gerado com: ' + rel['filtros_texto'])
        linhas = simpleSplit(texto, 'Helvetica', 8.5, largura - 2 * margem - 90)[:2]
        for i, parte in enumerate(linhas):
            self.drawString(margem, 1.3 * cm - i * 10, parte)
        self.drawRightString(largura - margem, 1.3 * cm, f'Página {self._pageNumber} de {total}')


def _estilos():
    base = ParagraphStyle('base', fontName='Helvetica', fontSize=8.5, leading=10.5)
    return {
        'celula': base,
        'cabeca': ParagraphStyle('cab', parent=base, fontName='Helvetica-Bold', textColor=colors.white),
        'titulo': ParagraphStyle('tit', parent=base, fontName='Helvetica-Bold', fontSize=12, leading=15,
                                 spaceBefore=10, spaceAfter=5, textColor=colors.HexColor(COR_ESCURA),
                                 keepWithNext=1),
        'nota': ParagraphStyle('nota', parent=base, textColor=colors.HexColor('#555555')),
    }


def _p(texto, estilo, vermelho=False):
    texto = escape(_latin(texto))
    if vermelho:
        texto = f'<font color="{COR_ATRASADA}">{texto}</font>'
    return Paragraph(texto, estilo)


def _tabela_pdf(estilos, cabecalho, linhas, larguras, vermelhas=()):
    dados = [[_p(c, estilos['cabeca']) for c in cabecalho]]
    for i, linha in enumerate(linhas):
        dados.append([_p(v, estilos['celula'], vermelho=(i in vermelhas)) for v in linha])
    estilo = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COR_ESCURA)),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#d1d5db')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]
    for i in range(1, len(dados)):
        if i % 2 == 0:
            estilo.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor(COR_ZEBRA)))
    tabela = Table(dados, colWidths=larguras, repeatRows=1)
    tabela.setStyle(TableStyle(estilo))
    return tabela


def gerar_pdf(rel):
    dados = rel['dados']
    est = _estilos()
    saida = io.BytesIO()
    doc = SimpleDocTemplate(
        saida, pagesize=landscape(A4),
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=2.5 * cm, bottomMargin=2.2 * cm,
        title=f"{rel['titulo']} - {rel['empresa']}", author=rel['empresa'],
    )
    util = landscape(A4)[0] - 3 * cm
    historia = []

    # Resumo
    historia.append(_p('Resumo', est['titulo']))
    linhas = _linhas_resumo(dados)
    atrasada = [i for i, l in enumerate(linhas) if l[0] == 'Atrasadas']
    historia.append(_tabela_pdf(est, ['Indicador', 'Valor', 'Observação'],
                                [(a, b, c) for a, b, c in linhas],
                                [util * 0.28, util * 0.14, util * 0.58], atrasada))

    # Distribuição por status e prioridade
    historia.append(_p('Distribuição', est['titulo']))
    t_status = _tabela_pdf(est, ['Status', 'Demandas'], [(x['nome'], x['qtd']) for x in dados['por_status']],
                           [util * 0.2, util * 0.1])
    t_prio = _tabela_pdf(est, ['Prioridade', 'Demandas'], [(x['nome'], x['qtd']) for x in dados['por_prioridade']],
                         [util * 0.2, util * 0.1])
    lado = Table([[t_status, t_prio]], colWidths=[util * 0.5, util * 0.5])
    lado.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 0)]))
    historia.append(lado)

    # Críticas
    historia.append(_p(f"Demandas críticas ({len(dados['criticas_lista'])})", est['titulo']))
    if dados['criticas_lista']:
        criticas = dados['criticas_lista']
        historia.append(_tabela_pdf(
            est, ['Nº', 'Título', 'Responsável', 'Status', 'Prazo'],
            [(c['id'], c['titulo'], c['responsavel'], c['status'],
              _data(c['prazo']) + (' (atrasada)' if c['atrasada'] else '')) for c in criticas],
            [util * 0.05, util * 0.45, util * 0.18, util * 0.14, util * 0.18],
            [i for i, c in enumerate(criticas) if c['atrasada']]))
    else:
        historia.append(_p('Nenhuma demanda crítica em aberto neste recorte.', est['nota']))

    # Por responsável
    historia.append(_p('Demandas em aberto por responsável', est['titulo']))
    historia.append(_tabela_pdf(est, ['Responsável', 'Em aberto'],
                                [(p['nome'], p['qtd']) for p in dados['por_responsavel']],
                                [util * 0.3, util * 0.12]))

    # Lista completa
    historia.append(_p(f"Demandas do relatório ({len(rel['demandas'])})", est['titulo']))
    if rel['demandas']:
        lista = rel['demandas']
        historia.append(_tabela_pdf(
            est, ['Nº', 'Título', 'Solicitante', 'Responsável', 'Prioridade', 'Status', 'Criação', 'Prazo', 'Conclusão'],
            [(d['id'], d['titulo'], d['solicitante'], d['responsavel'], d['prioridade'], d['status'],
              _data(d['criacao']), _data(d['prazo']) + (' (atrasada)' if d['atrasada'] else ''),
              _data(d['conclusao'])) for d in lista],
            [util * w for w in (0.04, 0.24, 0.11, 0.11, 0.08, 0.11, 0.08, 0.12, 0.11)],
            [i for i, d in enumerate(lista) if d['atrasada']]))
    else:
        historia.append(_p('Nenhuma demanda para os filtros escolhidos.', est['nota']))

    doc.build(historia, canvasmaker=lambda *a, **k: _CanvasNumerado(*a, relatorio=rel, **k))
    return saida.getvalue()


GERADORES = {'pdf': gerar_pdf, 'xlsx': gerar_xlsx, 'csv': gerar_csv}
