"""Testes automatizados do SGDI.

Rodar:   python -m unittest test_sgdi -v

Cada teste usa um banco temporário (variável SGDI_DB), então o seu
demandas.db real nunca é tocado.
"""
import csv
import io
import os
import re
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from datetime import date, datetime, timedelta

PASTA = os.path.dirname(os.path.abspath(__file__))
BANCO = os.path.join(tempfile.mkdtemp(prefix='sgdi_test_'), 'teste.db')
os.environ['SGDI_DB'] = BANCO
os.environ['NOME_EMPRESA'] = 'Acme Soluções & Cia'
sys.path.insert(0, PASTA)

import openpyxl  # noqa: E402
import pdfplumber  # noqa: E402

import app as sgdi  # noqa: E402


def recriar_banco():
    env = dict(os.environ, SGDI_DB=BANCO)
    subprocess.run([sys.executable, 'init_db.py'], cwd=PASTA, env=env, check=True, capture_output=True)


def sql(consulta, params=()):
    con = sqlite3.connect(BANCO)
    try:
        r = con.execute(consulta, params).fetchall()
        con.commit()
        return r
    finally:
        con.close()


def texto_pdf(conteudo):
    with pdfplumber.open(io.BytesIO(conteudo)) as pdf:
        paginas = [p.extract_text() or '' for p in pdf.pages]
    return paginas


def ler_csv(conteudo):
    texto = conteudo.decode('utf-8-sig')
    return list(csv.reader(io.StringIO(texto), delimiter=';'))


def linhas_csv_demandas(conteudo):
    """Linhas de dados (entre o cabeçalho de colunas e o rodapé)."""
    linhas = ler_csv(conteudo)
    i = next(k for k, l in enumerate(linhas) if l and l[0] == 'Nº')
    dados = []
    for l in linhas[i + 1:]:
        if not l:
            break
        dados.append(l)
    return dados


class Base(unittest.TestCase):
    def setUp(self):
        recriar_banco()
        self.c = sgdi.app.test_client()


# ------------------------------------------------------------------ painel

class TestPainel(Base):
    def test_numeros_do_painel_batem_com_o_banco(self):
        con = sqlite3.connect(BANCO)
        hoje = date.today().isoformat()
        esperado = {
            'total': con.execute('select count(*) from demandas').fetchone()[0],
            'abertas': con.execute("select count(*) from demandas where status in ('Aberta','Em andamento')").fetchone()[0],
            'concluidas': con.execute("select count(*) from demandas where status='Concluída'").fetchone()[0],
            'canceladas': con.execute("select count(*) from demandas where status='Cancelada'").fetchone()[0],
            'atrasadas': con.execute("select count(*) from demandas where status in ('Aberta','Em andamento') and prazo < ?", (hoje,)).fetchone()[0],
        }
        con.close()
        conn = sgdi.get_db()
        k = sgdi.calcular_painel(conn, [], [])['kpis']
        conn.close()
        for chave, valor in esperado.items():
            self.assertEqual(k[chave], valor, chave)
        self.assertEqual(k['abertas'] + k['concluidas'] + k['canceladas'], k['total'])

    def test_cancelada_nunca_e_atrasada_nem_critica(self):
        sql("UPDATE demandas SET status='Cancelada', prazo='2020-01-01', prioridade='Alta' WHERE id=2")
        conn = sgdi.get_db()
        d = sgdi.calcular_painel(conn, [], [])
        conn.close()
        self.assertNotIn(2, [c['id'] for c in d['criticas_lista']])
        atrasadas_banco = sql("select count(*) from demandas where status in ('Aberta','Em andamento') and prazo < ?",
                              (date.today().isoformat(),))[0][0]
        self.assertEqual(d['kpis']['atrasadas'], atrasadas_banco)

    def test_tempo_medio_ponderado_e_sem_canceladas(self):
        # concluídas do exemplo: dias 5, 3, 2 (peso 1) e 6 (Alta, peso 2)
        conn = sgdi.get_db()
        d = sgdi.calcular_painel(conn, [], [])
        conn.close()
        self.assertEqual(d['tempo_ponderado'], round((5 + 3 + 2 + 6 * 2) / 5, 1))
        self.assertEqual(d['tempo_simples'], 4.0)
        # uma cancelada com data_conclusao não pode entrar na média
        sql("UPDATE demandas SET data_conclusao='2024-02-20 10:00:00' WHERE status='Cancelada'")
        conn = sgdi.get_db()
        d2 = sgdi.calcular_painel(conn, [], [])
        conn.close()
        self.assertEqual(d2['tempo_ponderado'], d['tempo_ponderado'])
        self.assertEqual(d2['tempo_base'], 4)

    def test_paginas_abrem(self):
        for url in ['/', '/painel', '/painel?parcial=1', '/nova_demanda', '/editar/1', '/detalhes/1', '/usuarios']:
            self.assertEqual(self.c.get(url).status_code, 200, url)

    def test_atrasada_aparece_em_vermelho(self):
        html = self.c.get('/painel').get_data(as_text=True)
        self.assertIn('linha-atrasada', html)
        self.assertIn('texto-atrasada', html)
        self.assertIn('linha-atrasada', self.c.get('/').get_data(as_text=True))

    def test_grafico_de_evolucao_e_svg(self):
        html = self.c.get('/painel').get_data(as_text=True)
        self.assertIn('<svg', html)
        self.assertIn('<polyline', html)


class TestFiltros(Base):
    def total(self, qs):
        html = self.c.get('/painel?' + qs).get_data(as_text=True)
        return int(re.search(r'kpi-numero">(\d+)', html).group(1))

    def test_filtros_individuais_e_combinados_batem_com_sql(self):
        casos = {
            'prioridade=Alta': "prioridade='Alta'",
            'status=Cancelada': "status='Cancelada'",
            'responsavel_id=1': "responsavel_id=1",
            'de=2024-01-20&ate=2024-01-25': "substr(data_criacao,1,10) between '2024-01-20' and '2024-01-25'",
            'prioridade=Alta&status=Aberta': "prioridade='Alta' and status='Aberta'",
            'responsavel_id=2&prioridade=Alta&de=2024-01-15&ate=2024-01-16':
                "responsavel_id=2 and prioridade='Alta' and substr(data_criacao,1,10) between '2024-01-15' and '2024-01-16'",
            'de=2024-01-22': "substr(data_criacao,1,10) >= '2024-01-22'",
            'ate=2024-01-17': "substr(data_criacao,1,10) <= '2024-01-17'",
        }
        for qs, where in casos.items():
            esperado = sql(f'select count(*) from demandas where {where}')[0][0]
            self.assertEqual(self.total(qs), esperado, qs)

    def test_filtros_invalidos_sao_ignorados(self):
        total = sql('select count(*) from demandas')[0][0]
        self.assertEqual(self.total('de=lixo&ate=99&status=Xyz&prioridade=Zzz&responsavel_id=abc'), total)
        self.assertEqual(self.c.get('/painel?de=lixo&status=Xyz').status_code, 200)

    def test_filtro_vale_para_graficos_criticas_e_responsaveis(self):
        conn = sgdi.get_db()
        cond, params, filtros = sgdi._filtros_comuns({'prioridade': 'Baixa'})
        d = sgdi.calcular_painel(conn, cond, params)
        conn.close()
        self.assertEqual(d['criticas_lista'], [])  # Baixa não tem críticas
        self.assertEqual(sum(x['qtd'] for x in d['por_prioridade']), d['kpis']['total'])
        self.assertEqual(sum(x['qtd'] for x in d['por_status']), d['kpis']['total'])
        abertas_baixa = sql("select count(*) from demandas where prioridade='Baixa' and status in ('Aberta','Em andamento')")[0][0]
        self.assertEqual(sum(p['qtd'] for p in d['por_responsavel']), abertas_baixa)

    def test_lista_aceita_periodo(self):
        html = self.c.get('/?de=2024-01-20&ate=2024-01-25').get_data(as_text=True)
        esperado = sql("select count(*) from demandas where substr(data_criacao,1,10) between '2024-01-20' and '2024-01-25'")[0][0]
        self.assertIn(f'{esperado} demandas encontradas', html)


# -------------------------------------------------------- status e conclusão

class TestFluxoDeStatus(Base):
    def editar(self, id, status, prazo=''):
        return self.c.post(f'/editar/{id}', data=dict(
            titulo='t', descricao='d', solicitante_id=1, responsavel_id=2,
            prioridade='Média', status=status, prazo=prazo))

    def test_concluir_grava_data_e_reabrir_limpa(self):
        self.assertEqual(self.editar(2, 'Concluída').status_code, 302)
        status, data = sql('select status, data_conclusao from demandas where id=2')[0]
        self.assertEqual(status, 'Concluída')
        self.assertIsNotNone(data)
        self.editar(2, 'Aberta')
        self.assertIsNone(sql('select data_conclusao from demandas where id=2')[0][0])

    def test_cancelar_nao_grava_data_de_conclusao(self):
        self.editar(2, 'Cancelada')
        self.assertEqual(sql('select status, data_conclusao from demandas where id=2')[0], ('Cancelada', None))

    def test_concluida_mantem_data_original(self):
        original = sql('select data_conclusao from demandas where id=4')[0][0]
        self.editar(4, 'Concluída')
        self.assertEqual(sql('select data_conclusao from demandas where id=4')[0][0], original)

    def test_prazo_invalido_e_rejeitado(self):
        self.assertEqual(self.editar(2, 'Aberta', prazo='abc').status_code, 400)

    def test_nova_demanda_com_prazo(self):
        r = self.c.post('/nova_demanda', data=dict(
            titulo='Nova', descricao='d', solicitante_id=1, responsavel_id=1,
            prioridade='Alta', status='Aberta', prazo='2030-05-01'))
        self.assertEqual(r.status_code, 302)
        self.assertEqual(sql("select prazo from demandas where titulo='Nova'")[0][0], '2030-05-01')


# --------------------------------------------------------------- exportação

class TestExportacao(Base):
    FILTROS = 'prioridade=Alta&de=2024-01-01&ate=2024-03-31'
    EMPRESA = 'Acme Soluções & Cia'

    def baixar(self, formato, qs=''):
        r = self.c.get(f'/exportar/{formato}' + (('?' + qs) if qs else ''))
        self.assertEqual(r.status_code, 200, formato)
        return r

    # --- download
    def test_cabecalhos_http_e_nome_do_arquivo(self):
        esperado = {'pdf': 'application/pdf',
                    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    'csv': 'text/csv'}
        for formato, mime in esperado.items():
            r = self.baixar(formato)
            self.assertEqual(r.mimetype, mime)
            self.assertRegex(r.headers['Content-Disposition'],
                             rf'attachment; filename=relatorio-demandas-\d{{8}}-\d{{6}}\.{formato}')

    def test_formato_invalido_da_404(self):
        for f in ['doc', 'exe', 'PDF2', '..']:
            self.assertEqual(self.c.get(f'/exportar/{f}').status_code, 404, f)

    def test_botoes_de_exportar_na_tela(self):
        html = self.c.get('/painel?prioridade=Alta').get_data(as_text=True)
        for f in ('pdf', 'xlsx', 'csv'):
            self.assertIn(f'/exportar/{f}?prioridade=Alta', html)

    # --- pedido 02: cabeçalho com nome da empresa e data de geração
    # --- pedido 03: filtros aplicados no rodapé
    def test_pdf_tem_cabecalho_e_rodape_em_todas_as_paginas(self):
        paginas = texto_pdf(self.baixar('pdf', self.FILTROS).data)
        self.assertGreaterEqual(len(paginas), 1)
        for i, t in enumerate(paginas, start=1):
            self.assertIn(self.EMPRESA, t)
            self.assertRegex(t, r'Gerado em \d{2}/\d{2}/\d{4} às \d{2}:\d{2}')
            self.assertIn('Gerado com: Prioridade Alta · Jan–Mar 2024', t)
            self.assertIn(f'Página {i} de {len(paginas)}', t)

    def test_pdf_sem_filtro_diz_que_nao_ha_filtro(self):
        t = ' '.join(texto_pdf(self.baixar('pdf').data))
        self.assertIn('Gerado com: nenhum filtro (todas as demandas)', t)

    def test_data_de_geracao_e_a_de_agora(self):
        antes = datetime.now().replace(second=0, microsecond=0)
        t = texto_pdf(self.baixar('pdf').data)[0]
        m = re.search(r'Gerado em (\d{2}/\d{2}/\d{4}) às (\d{2}:\d{2})', t)
        gerado = datetime.strptime(f'{m.group(1)} {m.group(2)}', '%d/%m/%Y %H:%M')
        self.assertLessEqual(abs((gerado - antes).total_seconds()), 120)

    def test_pdf_numeros_batem_com_o_painel(self):
        t = ' '.join(texto_pdf(self.baixar('pdf').data))
        conn = sgdi.get_db()
        d = sgdi.calcular_painel(conn, [], [])
        conn.close()
        k = d['kpis']
        self.assertRegex(t, rf"Total de demandas\s+{k['total']}")
        self.assertRegex(t, rf"Abertas\s+{k['abertas']}")
        self.assertRegex(t, rf"Atrasadas\s+{k['atrasadas']}")
        self.assertIn(f"{str(d['tempo_ponderado']).replace('.', ',')} dias", t)
        self.assertIn(f"Demandas do relatório ({k['total']})", t)

    def test_xlsx_estrutura_cabecalho_rodape_e_dados(self):
        wb = openpyxl.load_workbook(io.BytesIO(self.baixar('xlsx', self.FILTROS).data))
        self.assertEqual(wb.sheetnames, ['Resumo', 'Demandas', 'Críticas', 'Por responsável'])
        for ws in wb.worksheets:
            self.assertEqual(ws['A1'].value, self.EMPRESA, ws.title)
            self.assertRegex(ws['A3'].value, r'Gerado em \d{2}/\d{2}/\d{4} às \d{2}:\d{2}')
            ultima = [c.value for c in ws['A'] if c.value][-1]
            self.assertEqual(ultima, 'Gerado com: Prioridade Alta · Jan–Mar 2024', ws.title)
            # também no cabeçalho/rodapé de impressão
            self.assertIn('Acme', ws.oddHeader.left.text)
            self.assertIn('Prioridade Alta', ws.oddFooter.left.text)
            self.assertEqual(ws.page_setup.orientation, 'landscape')
        ws = wb['Demandas']
        esperado = sql("select count(*) from demandas where prioridade='Alta' "
                       "and substr(data_criacao,1,10) between '2024-01-01' and '2024-03-31'")[0][0]
        ids = [r[0].value for r in ws.iter_rows(min_row=6, max_col=1) if isinstance(r[0].value, int)]
        self.assertEqual(len(ids), esperado)
        self.assertEqual(ws['B5'].value, 'Título')
        self.assertTrue(all(r[0].value == 'Alta' for r in ws.iter_rows(min_row=6, min_col=5, max_col=5) if r[0].value in ('Alta', 'Média', 'Baixa')))

    def test_xlsx_datas_sao_datas_de_verdade(self):
        ws = openpyxl.load_workbook(io.BytesIO(self.baixar('xlsx').data))['Demandas']
        self.assertIsInstance(ws['G6'].value, datetime)

    def test_csv_estrutura_e_dados(self):
        bruto = self.baixar('csv', self.FILTROS).data
        self.assertTrue(bruto.startswith(b'\xef\xbb\xbf'), 'precisa de BOM para o Excel abrir acentos')
        linhas = ler_csv(bruto)
        self.assertEqual(linhas[0][0], self.EMPRESA)
        self.assertRegex(linhas[2][0], r'Gerado em \d{2}/\d{2}/\d{4} às \d{2}:\d{2}')
        self.assertEqual(linhas[-1][0], 'Gerado com: Prioridade Alta · Jan–Mar 2024')
        dados = linhas_csv_demandas(bruto)
        esperado = sql("select count(*) from demandas where prioridade='Alta' "
                       "and substr(data_criacao,1,10) between '2024-01-01' and '2024-03-31'")[0][0]
        self.assertEqual(len(dados), esperado)
        self.assertTrue(all(len(l) == 10 for l in dados))
        self.assertTrue(all(l[4] == 'Alta' for l in dados))

    def test_os_tres_formatos_listam_as_mesmas_demandas(self):
        qs = 'status=Aberta&responsavel_id=1'
        esperado = {r[0] for r in sql("select id from demandas where status='Aberta' and responsavel_id=1")}
        csv_ids = {int(l[0]) for l in linhas_csv_demandas(self.baixar('csv', qs).data)}
        ws = openpyxl.load_workbook(io.BytesIO(self.baixar('xlsx', qs).data))['Demandas']
        xlsx_ids = {r[0].value for r in ws.iter_rows(min_row=6, max_col=1) if isinstance(r[0].value, int)}
        self.assertEqual(csv_ids, esperado)
        self.assertEqual(xlsx_ids, esperado)
        t = ' '.join(texto_pdf(self.baixar('pdf', qs).data))
        self.assertIn(f'Demandas do relatório ({len(esperado)})', t)

    def test_filtro_de_periodo_formata_bem(self):
        casos = {
            'de=2025-01-01&ate=2025-03-31': 'Jan–Mar 2025',
            'de=2025-02-01&ate=2025-02-28': 'Fev 2025',
            'de=2024-11-01&ate=2025-02-28': 'Nov/2024–Fev/2025',
            'de=2025-01-10&ate=2025-02-20': '10/01/2025 a 20/02/2025',
            'de=2025-01-10': 'a partir de 10/01/2025',
            'ate=2025-02-20': 'até 20/02/2025',
        }
        for qs, texto in casos.items():
            rel = sgdi.montar_relatorio(dict(p.split('=') for p in qs.split('&')))
            self.assertEqual(rel['filtros_texto'], texto, qs)
        rel = sgdi.montar_relatorio({'prioridade': 'Alta', 'status': 'Aberta', 'responsavel_id': '1'})
        self.assertEqual(rel['filtros_texto'], 'Prioridade Alta · Status Aberta · Responsável João Silva')

    # --- casos extremos
    def test_relatorio_vazio_funciona_nos_tres_formatos(self):
        qs = 'de=2000-01-01&ate=2000-01-31'
        t = ' '.join(texto_pdf(self.baixar('pdf', qs).data))
        self.assertIn('Nenhuma demanda para os filtros escolhidos', t)
        self.assertIn('Gerado com: Jan 2000', t)
        ws = openpyxl.load_workbook(io.BytesIO(self.baixar('xlsx', qs).data))['Demandas']
        self.assertEqual(ws['B5'].value, 'Título')
        self.assertEqual(linhas_csv_demandas(self.baixar('csv', qs).data), [])

    def test_muitas_demandas_paginam_com_cabecalho_e_rodape_em_todas(self):
        con = sqlite3.connect(BANCO)
        for i in range(150):
            con.execute("INSERT INTO demandas (titulo,descricao,solicitante_id,responsavel_id,data_criacao,prioridade,status,prazo)"
                        " VALUES (?,?,?,?,?,?,?,?)",
                        (f'Demanda em massa {i}', 'x', 1, 2, '2024-02-01 10:00:00', 'Média', 'Aberta', '2020-01-01'))
        con.commit()
        con.close()
        paginas = texto_pdf(self.baixar('pdf').data)
        self.assertGreater(len(paginas), 3)
        for i, t in enumerate(paginas, start=1):
            self.assertIn(self.EMPRESA, t)
            self.assertIn(f'Página {i} de {len(paginas)}', t)
            self.assertIn('Gerado com:', t)
        todas = ' '.join(paginas)
        self.assertIn('Demanda em massa 149', todas)
        self.assertIn('Demanda em massa 0 ', todas.replace('\n', ' ') + ' ')
        self.assertEqual(len(linhas_csv_demandas(self.baixar('csv').data)), 16 + 150)

    def test_titulo_com_formula_nao_executa_no_excel_nem_no_csv(self):
        malicioso = '=HYPERLINK("http://x","clique")'
        sql("UPDATE demandas SET titulo=? WHERE id=1", (malicioso,))
        ws = openpyxl.load_workbook(io.BytesIO(self.baixar('xlsx').data))['Demandas']
        celulas = [c for row in ws.iter_rows(min_row=6, min_col=2, max_col=2) for c in row if c.value and 'HYPERLINK' in str(c.value)]
        self.assertTrue(celulas)
        self.assertTrue(all(c.data_type == 's' for c in celulas), 'fórmula virou fórmula de verdade')
        linhas = linhas_csv_demandas(self.baixar('csv').data)
        achado = [l for l in linhas if 'HYPERLINK' in l[1]]
        self.assertTrue(achado)
        self.assertTrue(all(not l[1].startswith('=') for l in achado))

    def test_caracteres_especiais_e_xml_nao_quebram_o_pdf(self):
        sql("UPDATE demandas SET titulo=? WHERE id=1", ('Teste <b>negrito</b> & "aspas" ção ñ → 日本',))
        self.assertEqual(self.baixar('pdf').data[:5], b'%PDF-')
        t = ' '.join(texto_pdf(self.baixar('pdf').data))
        self.assertIn('<b>negrito</b> & "aspas" ção ñ', t)

    def test_titulo_muito_longo_nao_quebra(self):
        sql("UPDATE demandas SET titulo=? WHERE id=1", ('palavra ' * 80,))
        for f in ('pdf', 'xlsx', 'csv'):
            self.baixar(f)

    def test_nome_da_empresa_com_e_comercial_no_excel(self):
        ws = openpyxl.load_workbook(io.BytesIO(self.baixar('xlsx').data))['Resumo']
        self.assertEqual(ws['A1'].value, 'Acme Soluções & Cia')
        self.assertEqual(ws.oddHeader.left.text, 'Acme Soluções && Cia')  # '&' escapado p/ impressão


# ----------------------------------------------------------------- migração

# ------------------------------------------- visão críticas + atrasadas

class TestVisaoCriticasAtrasadas(Base):
    def setUp(self):
        super().setUp()
        ontem = (date.today() - timedelta(days=1)).isoformat()
        futuro = (date.today() + timedelta(days=30)).isoformat()
        # 1 crítica atrasada, 1 crítica no prazo, 1 atrasada que não é crítica,
        # 1 crítica cancelada com prazo vencido (não conta)
        sql("UPDATE demandas SET prioridade='Média', status='Concluída', prazo=?", (futuro,))
        sql("UPDATE demandas SET prioridade='Alta', status='Aberta', prazo=? WHERE id=1", (ontem,))
        sql("UPDATE demandas SET prioridade='Alta', status='Em andamento', prazo=? WHERE id=2", (futuro,))
        sql("UPDATE demandas SET prioridade='Baixa', status='Aberta', prazo=? WHERE id=3", (ontem,))
        sql("UPDATE demandas SET prioridade='Alta', status='Cancelada', prazo=? WHERE id=4", (ontem,))

    def test_painel_destaca_so_criticas_atrasadas(self):
        conn = sgdi.get_db()
        d = sgdi.calcular_painel(conn, [], [])
        conn.close()
        self.assertEqual([c['id'] for c in d['criticas_atrasadas']], [1])
        html = self.c.get('/painel').get_data(as_text=True)
        self.assertIn('Críticas atrasadas', html)

    def test_filtro_restringe_painel_e_lista(self):
        html = self.c.get('/painel?visao=criticas_atrasadas').get_data(as_text=True)
        self.assertIn('selected>Críticas + atrasadas', html)
        lista = self.c.get('/?visao=criticas_atrasadas').get_data(as_text=True)
        self.assertEqual(lista.count('/detalhes/'), 2)  # 1 linha: link do título + botão ver

    def test_exportacao_traz_so_esses_casos(self):
        qs = 'visao=criticas_atrasadas'
        csv_dados = linhas_csv_demandas(self.c.get(f'/exportar/csv?{qs}').data)
        self.assertEqual([l[0] for l in csv_dados], ['1'])
        rel = sgdi.montar_relatorio({'visao': 'criticas_atrasadas'})
        self.assertEqual(rel['titulo'], 'Relatório de Demandas Críticas Atrasadas')
        self.assertEqual(rel['filtros_texto'], 'Críticas atrasadas')
        self.assertEqual(len(rel['demandas']), 1)
        for formato in ('pdf', 'xlsx'):
            self.assertEqual(self.c.get(f'/exportar/{formato}?{qs}').status_code, 200)
        t = ' '.join(texto_pdf(self.c.get(f'/exportar/pdf?{qs}').data))
        self.assertIn('Demandas do relatório (1)', t)

    def test_valor_invalido_e_ignorado(self):
        rel = sgdi.montar_relatorio({'visao': 'qualquer'})
        self.assertEqual(rel['filtros_texto'], 'nenhum filtro (todas as demandas)')


class TestMigracao(unittest.TestCase):
    def test_banco_antigo_migra_sem_perder_dados(self):
        con = sqlite3.connect(BANCO)
        con.close()
        os.remove(BANCO)
        con = sqlite3.connect(BANCO)
        con.executescript('''
        CREATE TABLE usuarios (id INTEGER PRIMARY KEY AUTOINCREMENT, nome TEXT NOT NULL UNIQUE);
        CREATE TABLE demandas (id INTEGER PRIMARY KEY AUTOINCREMENT, titulo TEXT NOT NULL, descricao TEXT NOT NULL,
            solicitante_id INTEGER NOT NULL, responsavel_id INTEGER NOT NULL, data_criacao TEXT NOT NULL,
            prioridade TEXT NOT NULL DEFAULT 'Média' CHECK(prioridade IN ('Alta','Média','Baixa')),
            status TEXT NOT NULL DEFAULT 'Aberta' CHECK(status IN ('Aberta','Em andamento','Concluída')),
            FOREIGN KEY (solicitante_id) REFERENCES usuarios(id), FOREIGN KEY (responsavel_id) REFERENCES usuarios(id));
        CREATE TABLE comentarios (id INTEGER PRIMARY KEY AUTOINCREMENT, demanda_id INTEGER NOT NULL,
            comentario TEXT NOT NULL, autor TEXT NOT NULL, data TEXT NOT NULL,
            FOREIGN KEY (demanda_id) REFERENCES demandas(id) ON DELETE CASCADE);
        INSERT INTO usuarios(nome) VALUES ('A'),('B');
        INSERT INTO demandas(titulo,descricao,solicitante_id,responsavel_id,data_criacao,prioridade,status) VALUES
            ('t1','d',1,2,'2024-01-01 10:00:00','Alta','Aberta'),('t2','d',1,1,'2024-01-01 10:00:00','Baixa','Concluída');
        INSERT INTO comentarios(demanda_id,comentario,autor,data) VALUES (1,'oi','x','2024-01-02 10:00:00');
        ''')
        con.commit()
        con.close()

        sgdi.migrar_banco()
        sgdi.migrar_banco()  # rodar de novo não pode quebrar nem duplicar

        self.assertEqual(sql('select count(*) from demandas')[0][0], 2)
        self.assertEqual(sql('select count(*) from comentarios')[0][0], 1)
        self.assertEqual(sql('PRAGMA foreign_key_check'), [])
        cols = [c[1] for c in sql('PRAGMA table_info(demandas)')]
        self.assertIn('prazo', cols)
        self.assertIn('data_conclusao', cols)

        c = sgdi.app.test_client()
        r = c.post('/editar/1', data=dict(titulo='t', descricao='d', solicitante_id=1, responsavel_id=2,
                                          prioridade='Alta', status='Cancelada', prazo=''))
        self.assertEqual(r.status_code, 302)
        for url in ['/', '/painel', '/exportar/pdf', '/exportar/xlsx', '/exportar/csv']:
            self.assertEqual(c.get(url).status_code, 200, url)
        recriar_banco()


if __name__ == '__main__':
    unittest.main(verbosity=2)
